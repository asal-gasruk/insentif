#!/usr/bin/env python3
"""Regenerate src/data/parameter-targets-excel.json from Pencapaian Agustus Excel.

Extracts TARGET only (not ACTUAL) for salesman team sheets.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

try:
    import openpyxl
except ImportError as e:
    raise SystemExit("Install openpyxl: pip install openpyxl") from e

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "Pencapaian Agustus 2026 - Rev-2.xlsx"
OUT = ROOT / "src" / "data" / "parameter-targets-excel.json"

PERIOD = "2026-08"
TEAM_ID_RE = re.compile(r"^[A-Z]{2,4}-[A-Z0-9-]+$")

# sheet → list of (paramId, target_col)  — 1-based cols; ID col always 3
SIMPLE_SHEETS: dict[str, list[tuple[str, int]]] = {
    "Pencapaian Sales Value": [("allProduct", 6)],
    "Pencapaian Effective Call": [("ec", 6)],
    "Pencapaian New Open Outlet": [("noo", 6)],
    "Pencapaian Active Outlet": [("ao", 6)],
    "Pencapaian Item per Outlet": [("iptIpo", 6)],
    "Pencapaian Product Display": [("pd", 6)],
    "Pencapaian Sales Contract": [("contract", 6)],
    "Pencapaian Co-Branding": [("coBranding", 6)],
}


def is_team_code(val) -> bool:
    if not isinstance(val, str):
        return False
    s = val.strip()
    return bool(TEAM_ID_RE.match(s))


def num(val) -> float | None:
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        return float(val)
    return None


def parse_simple(ws, param_cols: list[tuple[str, int]]) -> list[dict]:
    rows: list[dict] = []
    for r in range(1, ws.max_row + 1):
        code = ws.cell(r, 3).value
        if not is_team_code(code):
            continue
        code = str(code).strip()
        name = ws.cell(r, 4).value
        for param_id, tgt_col in param_cols:
            tgt = num(ws.cell(r, tgt_col).value)
            if tgt is None:
                continue
            rows.append(
                {
                    "teamCode": code,
                    "salesman": str(name).strip() if name else None,
                    "paramId": param_id,
                    "period": PERIOD,
                    "target": tgt,
                    "sheet": ws.title,
                }
            )
    return rows


def parse_product_focus(ws) -> list[dict]:
    """VOL target col 7, RO target col 11."""
    rows: list[dict] = []
    for r in range(1, ws.max_row + 1):
        code = ws.cell(r, 3).value
        if not is_team_code(code):
            continue
        code = str(code).strip()
        name = ws.cell(r, 4).value
        for param_id, tgt_col in (("focusVol", 7), ("focusRO", 11)):
            tgt = num(ws.cell(r, tgt_col).value)
            if tgt is None:
                continue
            rows.append(
                {
                    "teamCode": code,
                    "salesman": str(name).strip() if name else None,
                    "paramId": param_id,
                    "period": PERIOD,
                    "target": tgt,
                    "sheet": ws.title,
                }
            )
    return rows


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    targets: list[dict] = []

    for sheet_name, param_cols in SIMPLE_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            print(f"WARN: missing sheet {sheet_name}")
            continue
        targets.extend(parse_simple(wb[sheet_name], param_cols))

    if "Pencapaian Product Focus" in wb.sheetnames:
        targets.extend(parse_product_focus(wb["Pencapaian Product Focus"]))

    # Dedupe by (teamCode, paramId, period) — last wins
    keyed: dict[tuple, dict] = {}
    for t in targets:
        keyed[(t["teamCode"], t["paramId"], t["period"])] = t
    unique = list(keyed.values())
    unique.sort(key=lambda x: (x["period"], x["teamCode"], x["paramId"]))

    out = {
        "source": "docs/Pencapaian Agustus 2026 - Rev-2.xlsx",
        "period": PERIOD,
        "targets": unique,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

    by_param: dict[str, int] = {}
    for t in unique:
        by_param[t["paramId"]] = by_param.get(t["paramId"], 0) + 1
    print(f"Wrote {OUT.relative_to(ROOT)}")
    print(f"targets={len(unique)} by_param={by_param}")


if __name__ == "__main__":
    main()
