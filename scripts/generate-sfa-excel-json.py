#!/usr/bin/env python3
"""Regenerate src/data/sfa-excel.json from docs Excel (Sales Team SFA + Pembagian Insentif).

Sales Team SFA is a vertical table; header row may repeat mid-sheet and column
order can change (Jawa Tengah puts NIK Karyawan in column G).
"""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError as e:
    raise SystemExit("Install openpyxl: pip install openpyxl") from e

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "Insentif Team - Data Karyawan SFA.xlsx"
OUT = ROOT / "src" / "data" / "sfa-excel.json"

SKIP_JABATAN_TEAM = {
    "SPV",
    "ASM",
    "ASPS",
    "ASPR",
    "HORECA SPV",
    "FSR",
    "Task Force",
    "KEY ACCOUNT SUPERVISOR",
    "Key Account Supervisor",
    "WAREHOUSE CABANG SUPERVISOR",
    "ASS INDIRECT",
    "ASPR INDIRECT",
    "ASS INDIRECT KEBUMEN",
    "SPV DIST",
}
# Fase berikutnya — bukan Master Tim sales
OUT_OF_SCOPE_JABATAN = {
    "Helper Gudang",
    "Delivery Man",
    "Helper Delivery",
    "Driver Distributor",
    "MIX",
}
LEAD_JABATAN = {
    "Salesman",
    "Canvas Motoris",
    "Sales Executive",
    "Horeca Executive",
    "Key Account Executive",
    "Sales TO",
    "Sales Exclusive",
}
MEMBER_JABATAN = {"Driver", "Helper"}

DEFAULT_COLMAP = {"nama": 3, "nik": 4, "kode": 5, "jabatan": 6, "mobil": 7}


def is_vacant(p: dict) -> bool:
    n = (p.get("nama") or "").lower()
    return "vacant" in n or not p.get("nama")


def norm_nik(nik) -> str | None:
    if nik is None:
        return None
    s = str(nik).strip()
    if s in ("", "0"):
        return None
    # NIK harus numerik; hindari kode tim yang terselip di kolom salah
    if not re.fullmatch(r"\d+", s):
        return None
    return s


def is_header_row(vals: list) -> bool:
    labels = {str(v).strip() for v in vals if v is not None and str(v).strip()}
    return "AREA" in labels and "CABANG" in labels and "NAMA" in labels


def colmap_from_header(ws, row: int) -> dict[str, int]:
    headers = {
        str(ws.cell(row, c).value).strip(): c
        for c in range(1, 8)
        if ws.cell(row, c).value is not None
    }
    return {
        "nama": headers.get("NAMA", DEFAULT_COLMAP["nama"]),
        "nik": headers.get("NIK Karyawan", DEFAULT_COLMAP["nik"]),
        "kode": headers.get("Kode", DEFAULT_COLMAP["kode"]),
        "jabatan": headers.get("JABATAN", DEFAULT_COLMAP["jabatan"]),
        "mobil": headers.get("Jenis Mobil", DEFAULT_COLMAP["mobil"]),
    }


def should_skip_jabatan_as_team_break(j: str) -> bool:
    if not j:
        return False
    if j in SKIP_JABATAN_TEAM or j in OUT_OF_SCOPE_JABATAN:
        return True
    if j.startswith("SP-"):
        return True
    if "INDIRECT" in j.upper() or j.endswith("DIST"):
        return True
    return False


def parse_sfa(ws) -> list[dict]:
    """Parse vertical Sales Team SFA with mid-sheet header / column-order changes."""
    people: list[dict] = []
    area = cabang = None
    colmap = dict(DEFAULT_COLMAP)

    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 8)]
        if is_header_row(vals):
            colmap = colmap_from_header(ws, r)
            continue

        a, cab = vals[0], vals[1]
        if a is not None and str(a).strip() and str(a).strip().upper() != "AREA":
            area = str(a).strip()
        if cab is not None and str(cab).strip():
            cab_s = str(cab).strip()
            if cab_s.startswith("*"):
                continue
            if cab_s.upper() != "CABANG":
                cabang = cab_s

        nama = ws.cell(r, colmap["nama"]).value
        nik = ws.cell(r, colmap["nik"]).value
        kode = ws.cell(r, colmap["kode"]).value
        jabatan = ws.cell(r, colmap["jabatan"]).value
        mobil = ws.cell(r, colmap["mobil"]).value

        if not nama and not jabatan and not kode and not nik:
            continue
        if isinstance(nama, str) and nama.strip().upper() == "NAMA":
            continue

        people.append(
            {
                "area": area,
                "cabang": cabang,
                "nama": str(nama).strip() if nama else None,
                "nik": norm_nik(nik),
                "kode": None
                if kode is None or str(kode).strip() in ("", "0")
                else str(kode).strip(),
                "jabatan": str(jabatan).strip() if jabatan else None,
                "mobil": str(mobil).strip() if mobil else None,
            }
        )
    return people


def extract_teams(people_list: list[dict]) -> list[dict]:
    teams: list[dict] = []
    current = None
    for p in people_list:
        j = (p.get("jabatan") or "").strip()
        if not j:
            continue
        if should_skip_jabatan_as_team_break(j):
            if current:
                teams.append(current)
                current = None
            continue
        if j in LEAD_JABATAN:
            if current:
                teams.append(current)
            current = {
                "lead": p,
                "members": [p],
                "kode": p.get("kode"),
                "cabang": p.get("cabang"),
                "area": p.get("area"),
            }
            continue
        if current is not None and j in MEMBER_JABATAN:
            current["members"].append(p)
            continue
        if current:
            teams.append(current)
            current = None
    if current:
        teams.append(current)
    return teams


def resolve_team_role(kode: str | None, lead_j: str) -> tuple[str, str]:
    kode_u = (kode or "").upper()
    if "HRC" in kode_u or lead_j == "Horeca Executive":
        return "Horeca", "sales-horeca"
    if "MT" in kode_u or lead_j in ("Sales Executive", "Key Account Executive"):
        return "MT", "sales-mt"
    # Sales TO / Sales Exclusive / Canvas Motoris → GT canvasser (kode *-GT-*)
    return "GT", "canvasser"


def is_individu_jabatan(j: str) -> bool:
    if not j:
        return False
    if j in SKIP_JABATAN_TEAM or j.startswith("SP-"):
        return True
    if "INDIRECT" in j.upper() or j.endswith("DIST"):
        return True
    return False


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    people = parse_sfa(wb["Sales Team SFA"])

    by_cabang: dict[tuple, list] = defaultdict(list)
    for p in people:
        if p["cabang"]:
            by_cabang[(p["area"], p["cabang"])].append(p)

    raw_teams = []
    for plist in by_cabang.values():
        for t in extract_teams(plist):
            if is_vacant(t["lead"]):
                continue
            lead_j = t["lead"]["jabatan"]
            if lead_j not in LEAD_JABATAN:
                continue
            raw_teams.append(t)

    ws2 = wb["Pembagian Insentif"]
    examples = []
    for r in range(9, 15):
        examples.append(
            {
                "kode": ws2.cell(r, 1).value,
                "salesman": ws2.cell(r, 2).value,
                "typeSales": ws2.cell(r, 6).value,
                "incentiveGross": ws2.cell(r, 7).value,
                "salesShare": ws2.cell(r, 8).value,
                "driverHelperShare": ws2.cell(r, 9).value,
            }
        )

    emp_map: dict[str, dict] = {}
    out_teams = []
    for t in raw_teams:
        members_out = []
        helper_idx = 0
        for m in t["members"]:
            if is_vacant(m) or not m["nik"]:
                continue
            j = m["jabatan"]
            if j in LEAD_JABATAN:
                pos = "salesman"
            elif j == "Driver":
                pos = "driver"
            elif j == "Helper":
                helper_idx += 1
                pos = "helper1" if helper_idx == 1 else "helper2"
            else:
                continue
            emp_map[m["nik"]] = {
                "nik": m["nik"],
                "name": m["nama"],
                "cabang": m["cabang"] or t["cabang"],
                "area": m["area"] or t["area"],
                "jabatan": j,
                "kode": m.get("kode"),
                "mobil": m.get("mobil"),
            }
            members_out.append({"nik": m["nik"], "position": pos, "jabatan": j})
        if not members_out or not any(x["position"] == "salesman" for x in members_out):
            continue
        kode = t.get("kode") or f"TEAM-{t['lead']['nik']}"
        team_type, role = resolve_team_role(kode, t["lead"]["jabatan"])
        out_teams.append(
            {
                "code": kode,
                "name": f"{kode} · {t['lead']['nama']}",
                "teamType": team_type,
                "roleId": role,
                "cabang": t["cabang"],
                "area": t["area"],
                "members": members_out,
            }
        )

    individu = []
    seen_ind = set()
    for p in people:
        j = (p.get("jabatan") or "").strip()
        if is_vacant(p) or not p.get("nik"):
            continue
        if is_individu_jabatan(j):
            if p["nik"] in seen_ind:
                continue
            seen_ind.add(p["nik"])
            individu.append(
                {
                    "nik": p["nik"],
                    "name": p["nama"],
                    "cabang": p["cabang"],
                    "area": p["area"],
                    "jabatan": j,
                    "kode": p.get("kode"),
                }
            )

    cabangs = sorted(
        {
            (p["area"], p["cabang"])
            for p in people
            if p["cabang"] and not str(p["cabang"]).startswith("*")
        }
    )

    out = {
        "source": "docs/Insentif Team - Data Karyawan SFA.xlsx",
        "sheets": {"splits": "Pembagian Insentif", "teams": "Sales Team SFA"},
        "poolRule": "non-salesman share remaining pool equally",
        "splits": {
            "3": {"salesman": 0.55, "poolNonSalesman": 0.45},
            "2": {"salesman": 0.7, "poolNonSalesman": 0.3},
            "1": {"salesman": 1.0, "poolNonSalesman": 0.0},
        },
        "examples": examples,
        "cabangs": [{"area": a, "cabang": c} for a, c in cabangs],
        "employees": list(emp_map.values()),
        "individu": individu,
        "teams": out_teams,
    }

    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    sizes = Counter(len(t["members"]) for t in out_teams)
    by_area = Counter(t["area"] for t in out_teams)
    print(f"Wrote {OUT.relative_to(ROOT)}")
    print(
        f"teams={len(out_teams)} employees={len(emp_map)} "
        f"individu={len(individu)} sizes={dict(sizes)}"
    )
    print(f"by_area={dict(by_area)}")


if __name__ == "__main__":
    main()
